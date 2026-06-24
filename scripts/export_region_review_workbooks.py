from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
REPORT_DATE = "2026-06-15"
SCHEMA_VERSION = f"{REPORT_DATE}.review"
REGION_IDS = ("trondelag", "bornholm")

REGION_CATALOG_XLSX = ROOT / "docs" / f"REGION_CATALOG_SIDE_BY_SIDE_{REPORT_DATE}.xlsx"
PARAMETER_CATALOG_XLSX = ROOT / "docs" / f"REGION_PARAMETER_BUFFER_CATALOG_{REPORT_DATE}.xlsx"


HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
README_FILL = PatternFill("solid", fgColor="E2F0F3")
THIN_BORDER = Border(bottom=Side(style="thin", color="D9E2E5"))


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise ValueError(f"Expected object in {path}")
    return data


def save_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(data, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def rel(path: Path | str | None) -> str:
    if path is None:
        return ""
    raw = str(path).replace("\\", "/")
    return raw


def repo_path(path: str | None) -> Path | None:
    if not path:
        return None
    raw = Path(path)
    return raw if raw.is_absolute() else ROOT / raw


def exists_text(path: str | None) -> str:
    p = repo_path(path)
    if not p:
        return ""
    return "yes" if p.exists() else "missing"


def as_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def join_values(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def make_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def add_table_sheet(
    wb: Workbook,
    title: str,
    rows: list[dict[str, Any]],
    headers: list[str] | None = None,
    widths: dict[str, int] | None = None,
) -> None:
    ws = wb.create_sheet(title)
    headers = headers or sorted({key for row in rows for key in row})
    ws.append(headers)
    for row in rows:
        ws.append([as_cell(row.get(header, "")) for header in headers])

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = THIN_BORDER

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = THIN_BORDER

    ws.freeze_panes = "A2"
    if ws.max_row and ws.max_column:
        ws.auto_filter.ref = ws.dimensions

    widths = widths or {}
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        max_len = len(str(header))
        for cell in ws[letter]:
            text = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(max(len(part) for part in text.splitlines() or [""]), 90))
        ws.column_dimensions[letter].width = widths.get(header, min(max(max_len + 2, 12), 70))


def add_readme_sheet(wb: Workbook, title: str, rows: list[tuple[str, Any]]) -> None:
    ws = wb.create_sheet(title, 0)
    ws.append(["Field", "Value"])
    for key, value in rows:
        ws.append([key, as_cell(value)])
    for cell in ws[1]:
        cell.fill = README_FILL
        cell.font = Font(bold=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = THIN_BORDER
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100


def load_region_context() -> dict[str, dict[str, Any]]:
    contexts: dict[str, dict[str, Any]] = {}
    for region_id in REGION_IDS:
        region_path = ROOT / "regions" / region_id / "region.json"
        region = load_json(region_path)
        registry_path = repo_path((region.get("manifests") or {}).get("acceptance_registry"))
        registry = load_json(registry_path) if registry_path and registry_path.exists() else {}
        contexts[region_id] = {
            "region": region,
            "region_path": region_path,
            "registry": registry,
            "registry_path": registry_path,
        }
    return contexts


def group_readiness(region_id: str, group_id: str, layer_count: int) -> str:
    if region_id == "trondelag" and group_id in {"aviation_approach", "aviation_bird", "military", "coastal"}:
        return "not_wired_for_trondelag_prototype"
    if region_id == "trondelag" and group_id == "settlement":
        return "wired_proxy_review"
    if region_id == "trondelag" and layer_count > 0:
        return "wired_prototype_review"
    if region_id == "bornholm":
        return "source_registry_review_r9_runtime_not_rebuilt"
    return "registry_only_review"


def layer_readiness(region_id: str, group_id: str, layer_id: str) -> str:
    if region_id == "trondelag" and group_id in {"aviation_approach", "aviation_bird", "military", "coastal"}:
        return "not_wired_for_trondelag_prototype"
    if region_id == "trondelag" and layer_id == "population_points":
        return "proxy_250m_grid_review"
    if region_id == "bornholm":
        return "registry_source_review_r9_runtime_not_rebuilt"
    return "prototype_source_review"


def region_notes(region_id: str) -> list[str]:
    if region_id == "trondelag":
        return [
            "Do all distance, buffer and area geometry in EPSG:25832.",
            "Population and settlement buffers must be dissolved polygon buffers from the 250 m population-grid proxy.",
            "Do not expose Trondelag R8/R9 in the interactive app.",
            "Aviation, military and coastal groups are registry placeholders until Trondelag sources are wired.",
        ]
    return [
        "Do all distance, buffer and area geometry in EPSG:25833.",
        "Bornholm R9 data package is imported for V2 review, but the full parameter/buffer runtime has not been rebuilt as a clean R9-native workflow.",
        "R10 remains only as source/provenance in imported manifests and v1 runtime outputs.",
    ]


def runtime_rendering(region_id: str) -> dict[str, Any]:
    if region_id != "trondelag":
        return {}
    return {
        "population_buffer": {
            "render_mode": "dissolved_polygon_proxy",
            "source_layer_id": "population_points",
            "proxy_resolution_m": 250,
            "source_rds": "docs/geocontext/acceptance_framework/data/trondelag_prototype_assets/analysis_rds/population_points.rds",
            "render_script": "script/acceptance/render_trondelag_population_buffer.R",
            "cache_dir": "docs/geocontext/acceptance_framework/data/trondelag_prototype_assets/runtime_buffers",
            "cache_filename_template": "population_points_buffer_{buffer_m}m.geojson",
            "native_crs": "EPSG:25832",
            "web_crs": "EPSG:4326",
            "runtime_status": "proxy_250m_grid_dissolved_polygon",
            "user_facing_note": "Trondelag population uses dissolved 250 m grid-cell proxy polygons derived from centroids, not individual population points.",
        }
    }


def build_group_records(contexts: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for region_id, ctx in contexts.items():
        registry = ctx["registry"]
        layers_by_group: dict[str, list[dict[str, Any]]] = {}
        for layer in registry.get("layers") or []:
            layers_by_group.setdefault(str(layer.get("group_id")), []).append(layer)
        for group in registry.get("groups") or []:
            group_id = str(group.get("id"))
            layers = layers_by_group.get(group_id, [])
            records.append(
                {
                    "region_id": region_id,
                    "group_id": group_id,
                    "group_label": group.get("label", ""),
                    "analysis_kind": group.get("analysis_kind", ""),
                    "analysis_label": group.get("analysis_label", ""),
                    "default_m": group.get("analysis_default_m", ""),
                    "min_m": group.get("analysis_min_m", ""),
                    "max_m": group.get("analysis_max_m", ""),
                    "step_m": group.get("analysis_step_m", ""),
                    "blend_default": group.get("blend_default", ""),
                    "source_layer_count": len(layers),
                    "source_layer_ids": [layer.get("id") for layer in layers],
                    "readiness": group_readiness(region_id, group_id, len(layers)),
                    "source_registry": rel((ctx["region"].get("manifests") or {}).get("acceptance_registry")),
                    "interpretation": group.get("interpretation", ""),
                }
            )
    return records


def build_layer_records(contexts: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for region_id, ctx in contexts.items():
        for layer in ctx["registry"].get("layers") or []:
            group_id = str(layer.get("group_id", ""))
            records.append(
                {
                    "region_id": region_id,
                    "group_id": group_id,
                    "layer_id": layer.get("id", ""),
                    "label": layer.get("label", ""),
                    "layer_key": layer.get("layer_key", ""),
                    "filter_field": layer.get("filter_field", ""),
                    "filter_value": layer.get("filter_value", ""),
                    "filter_mode": layer.get("filter_mode", ""),
                    "readiness": layer_readiness(region_id, group_id, str(layer.get("id", ""))),
                    "note": layer.get("note", ""),
                }
            )
    return records


def wind_app_defaults() -> list[dict[str, Any]]:
    builder = [
        ("settlement_distance_m", "settlement", "Minimum distance to residents", 100, 100, 3000, 50, "distance_conflict"),
        ("road_distance_m", "transport", "Minimum distance to roads", 100, 50, 2000, 25, "distance_conflict"),
        ("grid_max_distance_m", "electrical", "Maximum distance to electrical infrastructure", 2000, 500, 15000, 250, "proximity_feasibility"),
        ("protected_buffer_m", "protected", "Buffer protected areas", 0, 0, 2000, 50, "hard_exclusion"),
        ("coastal_buffer_m", "coastal", "Buffer coastal/strand protection", 0, 0, 1000, 50, "hard_exclusion"),
        ("culture_buffer_m", "culture", "Buffer cultural environments", 0, 0, 1500, 50, "hard_exclusion"),
        ("reindeer_buffer_m", "reindeer", "Buffer reindeer husbandry", 0, 0, 5000, 100, "hard_exclusion"),
        ("aviation_approach_buffer_m", "aviation_approach", "Buffer aviation approach zones", 0, 0, 3000, 100, "hard_exclusion"),
        ("aviation_bird_distance_m", "aviation_bird", "Minimum distance aviation bird-collision sensitivity", 0, 0, 4000, 100, "distance_conflict"),
        ("military_buffer_m", "military", "Buffer military areas", 0, 0, 2000, 50, "hard_exclusion"),
        ("landscape_sensitivity_percent", "landscape", "Landscape sensitivity weight", 60, 0, 120, 5, "score_weight"),
    ]
    reference_overrides = {
        "settlement_distance_m": 500,
        "road_distance_m": 300,
        "grid_max_distance_m": 1000,
        "protected_buffer_m": 250,
        "culture_buffer_m": 100,
        "reindeer_buffer_m": 100,
    }
    rows: list[dict[str, Any]] = []
    for control_id, group_id, label, default_m, min_m, max_m, step_m, operation in builder:
        rows.append(
            {
                "technology": "wind",
                "default_profile": "builder_default",
                "control_id": control_id,
                "group_id": group_id,
                "label": label,
                "default_value": default_m,
                "min_value": min_m,
                "max_value": max_m,
                "step": step_m,
                "operation": operation,
                "source": "potential_app.py::_default_wind_params",
                "notes": "Current neutral builder default.",
            }
        )
        if control_id in reference_overrides:
            rows.append(
                {
                    "technology": "wind",
                    "default_profile": "reference_default",
                    "control_id": control_id,
                    "group_id": group_id,
                    "label": label,
                    "default_value": reference_overrides[control_id],
                    "min_value": min_m,
                    "max_value": max_m,
                    "step": step_m,
                    "operation": operation,
                    "source": "potential_app.py::_reference_default_wind_params",
                    "notes": "Reference scenario default used as a stronger planning preset.",
                }
            )
    return rows


def solar_app_defaults() -> list[dict[str, Any]]:
    return [
        {
            "technology": "solar",
            "default_profile": "applied_default",
            "control_id": "large_population_active",
            "group_id": "settlement",
            "label": "Large-scale solar population clearance",
            "default_active": True,
            "default_layer_ids": "region settlement defaults",
            "default_value": 500,
            "spec_default_value": "",
            "min_value": 100,
            "max_value": 3000,
            "step": 50,
            "operation": "distance_conflict",
            "source": "potential_app.py::DEFAULT_SOLAR_APPLIED_CONFIG",
            "notes": "Trondelag uses the 250 m grid-cell proxy; Bornholm still needs R9 runtime review.",
        },
        {
            "technology": "solar",
            "default_profile": "applied_default",
            "control_id": "large_protected_active",
            "group_id": "protected",
            "label": "Protected nature filter",
            "default_active": True,
            "default_layer_ids": ["protected_areas"],
            "default_value": 250,
            "spec_default_value": 0,
            "min_value": 0,
            "max_value": 2000,
            "step": 50,
            "operation": "hard_exclusion",
            "source": "potential_app.py::DEFAULT_SOLAR_APPLIED_CONFIG + SOLAR_FILTER_GROUP_SPECS",
            "notes": "App default is stricter than the group registry default.",
        },
        {
            "technology": "solar",
            "default_profile": "applied_default",
            "control_id": "large_land_use_active",
            "group_id": "land_use",
            "label": "Forest/land-use filter",
            "default_active": False,
            "default_layer_ids": ["forest_land_cover"],
            "default_value": 0,
            "spec_default_value": 0,
            "min_value": 0,
            "max_value": 1000,
            "step": 50,
            "operation": "hard_exclusion",
            "source": "potential_app.py::DEFAULT_SOLAR_APPLIED_CONFIG + SOLAR_FILTER_GROUP_SPECS",
            "notes": "Currently meaningful for Trondelag land-use controller; Bornholm has no equivalent registry group.",
        },
        {
            "technology": "solar",
            "default_profile": "applied_default",
            "control_id": "large_road_active",
            "group_id": "transport",
            "label": "Road clearance",
            "default_active": True,
            "default_layer_ids": "region transport defaults",
            "default_value": 300,
            "spec_default_value": 100,
            "min_value": 0,
            "max_value": 500,
            "step": 25,
            "operation": "distance_conflict",
            "source": "potential_app.py::DEFAULT_SOLAR_APPLIED_CONFIG + SOLAR_FILTER_GROUP_SPECS",
            "notes": "App default is stricter than the group registry default.",
        },
        {
            "technology": "solar",
            "default_profile": "applied_default",
            "control_id": "large_electrical_active",
            "group_id": "electrical",
            "label": "Near-grid feasibility",
            "default_active": False,
            "default_layer_ids": ["high_voltage_lines", "underground_cables"],
            "default_value": 2000,
            "spec_default_value": 2000,
            "min_value": 500,
            "max_value": 15000,
            "step": 250,
            "operation": "proximity_feasibility",
            "source": "potential_app.py::DEFAULT_SOLAR_APPLIED_CONFIG + SOLAR_FILTER_GROUP_SPECS",
            "notes": "Distance is only applied when the group is activated.",
        },
        {
            "technology": "solar",
            "default_profile": "applied_default",
            "control_id": "large_culture_active",
            "group_id": "culture",
            "label": "Cultural environment filter",
            "default_active": True,
            "default_layer_ids": ["cultural_preservation", "valuable_cultural_environment"],
            "default_value": 100,
            "spec_default_value": 0,
            "min_value": 0,
            "max_value": 1500,
            "step": 50,
            "operation": "hard_exclusion",
            "source": "potential_app.py::DEFAULT_SOLAR_APPLIED_CONFIG + SOLAR_FILTER_GROUP_SPECS",
            "notes": "App default is stricter than the group registry default.",
        },
        {
            "technology": "solar",
            "default_profile": "applied_default",
            "control_id": "large_reindeer_active",
            "group_id": "reindeer",
            "label": "Reindeer husbandry filter",
            "default_active": True,
            "default_layer_ids": ["reindeer_grazing_merged"],
            "default_value": 100,
            "spec_default_value": 0,
            "min_value": 0,
            "max_value": 5000,
            "step": 100,
            "operation": "hard_exclusion",
            "source": "potential_app.py::DEFAULT_SOLAR_APPLIED_CONFIG + SOLAR_FILTER_GROUP_SPECS",
            "notes": "Trondelag-specific. No Bornholm equivalent should be forced.",
        },
        {
            "technology": "solar",
            "default_profile": "applied_default",
            "control_id": "large_coastal_active",
            "group_id": "coastal",
            "label": "Coastal/strand protection filter",
            "default_active": False,
            "default_layer_ids": [],
            "default_value": 0,
            "spec_default_value": 0,
            "min_value": 0,
            "max_value": 1000,
            "step": 50,
            "operation": "hard_exclusion",
            "source": "potential_app.py::DEFAULT_SOLAR_APPLIED_CONFIG + SOLAR_FILTER_GROUP_SPECS",
            "notes": "Available as a cautious coastal filter when sources are wired.",
        },
    ]


def build_region_catalogs(contexts: dict[str, dict[str, Any]]) -> dict[str, Path]:
    group_records = build_group_records(contexts)
    layer_records = build_layer_records(contexts)
    groups_by_region: dict[str, list[dict[str, Any]]] = {region_id: [] for region_id in REGION_IDS}
    layers_by_region_group: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for layer in layer_records:
        layers_by_region_group.setdefault((str(layer["region_id"]), str(layer["group_id"])), []).append(layer)

    for group in group_records:
        group_copy = deepcopy(group)
        group_copy["layers"] = layers_by_region_group.get((str(group["region_id"]), str(group["group_id"])), [])
        groups_by_region[str(group["region_id"])].append(group_copy)

    output_paths: dict[str, Path] = {}
    for region_id, ctx in contexts.items():
        region = ctx["region"]
        runtime_rendering_config = runtime_rendering(region_id)
        catalog = {
            "schema_version": SCHEMA_VERSION,
            "created_at": REPORT_DATE,
            "status": "review_catalog_not_runtime_contract",
            "region_id": region_id,
            "display_name": region.get("display_name", region_id),
            "native_crs": region.get("native_crs") or (region.get("crs") or {}).get("native"),
            "web_crs": region.get("web_crs") or (region.get("crs") or {}).get("web"),
            "source_region_package": rel(ctx["region_path"].relative_to(ROOT)),
            "source_registry": rel((region.get("manifests") or {}).get("acceptance_registry")),
            "notes": region_notes(region_id),
            "groups": groups_by_region.get(region_id, []),
            "app_defaults": {
                "solar": solar_app_defaults(),
                "wind": wind_app_defaults(),
            },
            "review_checklist": [
                "Confirm that every selected source layer exists in the regional data package.",
                "Confirm CRS before any distance, buffer or area operation.",
                "Confirm whether buffer output is polygon, H3 coverage share or distance table.",
                "Confirm whether the control is active by default in the app or only available as an advanced option.",
                "For Bornholm, decide whether to rebuild true R9 PEY runtime or label the current mode as compatibility.",
                "For Trondelag, keep population buffers as dissolved polygon buffers from the 250 m proxy.",
            ],
        }
        if runtime_rendering_config:
            catalog = {**catalog, "runtime_rendering": runtime_rendering_config}
            catalog = {
                key: catalog[key]
                for key in [
                    "schema_version",
                    "created_at",
                    "status",
                    "region_id",
                    "display_name",
                    "native_crs",
                    "web_crs",
                    "source_region_package",
                    "source_registry",
                    "notes",
                    "runtime_rendering",
                    "groups",
                    "app_defaults",
                    "review_checklist",
                ]
            }
        path = ROOT / "regions" / region_id / "parameter_buffers.json"
        save_json(path, catalog)
        output_paths[region_id] = path
    return output_paths


def manifest_summary(region: dict[str, Any], kind: str, path: str | None) -> dict[str, Any]:
    p = repo_path(path)
    data: dict[str, Any] = {}
    if p and p.exists() and p.suffix.lower() == ".json":
        data = load_json(p)
    return {
        "region_id": region.get("region_id"),
        "manifest_type": kind,
        "path": rel(path),
        "exists": exists_text(path),
        "status": data.get("status", data.get("data_status", "")),
        "id": data.get("analysis_id", data.get("dataset_id", data.get("model_id", data.get("analysis_id", "")))),
        "h3_resolution": data.get("h3_resolution", data.get("source_h3_resolution", data.get("default_h3_resolution", ""))),
        "available_h3": data.get("available_h3_resolutions", ""),
        "rows": data.get("rows", ""),
        "source_h3_resolution": data.get("source_h3_resolution", ""),
        "note": data.get("note", data.get("method_note", data.get("runtime_intent", ""))),
    }


def export_region_workbook(contexts: dict[str, dict[str, Any]], catalog_paths: dict[str, Path]) -> None:
    wb = make_workbook()
    add_readme_sheet(
        wb,
        "README",
        [
            ("Purpose", "Human review workbook for comparing active V2 region packages."),
            ("Generated", REPORT_DATE),
            ("Regions", ", ".join(REGION_IDS)),
            ("Markdown source", f"docs/REGION_CATALOG_SIDE_BY_SIDE_{REPORT_DATE}.md"),
            ("Parameter workbook", rel(PARAMETER_CATALOG_XLSX.relative_to(ROOT))),
            ("Important caveat", "This workbook is a review aid. The JSON manifests remain the runtime sources."),
        ],
    )

    primary_rows = []
    for region_id, ctx in contexts.items():
        region = ctx["region"]
        manifests = region.get("manifests") or {}
        primary_rows.append(
            {
                "region_id": region_id,
                "region_package": rel(ctx["region_path"].relative_to(ROOT)),
                "parameter_buffer_catalog": rel(catalog_paths[region_id].relative_to(ROOT)),
                "landscape_manifest": manifests.get("landscape", region.get("landscape_manifest", "")),
                "potential_manifest": manifests.get("potential", region.get("potential_manifest", "")),
                "social_acceptance_manifest": manifests.get("social_acceptance", region.get("social_acceptance_manifest", "")),
                "scenario_manifest": manifests.get("scenarios", region.get("scenario_manifest", "")),
                "acceptance_registry": manifests.get("acceptance_registry", ""),
                "handoff": (region.get("source_package") or {}).get("handoff", "docs/TRONDELAG_PDF_LANDSCAPE_HANDOFF_2026-05-13.md" if region_id == "trondelag" else ""),
            }
        )
    add_table_sheet(wb, "Primary Files", primary_rows)

    contract_rows = []
    for region_id, ctx in contexts.items():
        region = ctx["region"]
        landing = region.get("landing_card") or {}
        h3 = region.get("h3") or {}
        contract_rows.append(
            {
                "region_id": region_id,
                "display_name": region.get("display_name", ""),
                "country": region.get("country", ""),
                "status": region.get("status", ""),
                "data_status": region.get("data_status", ""),
                "native_crs": region.get("native_crs") or (region.get("crs") or {}).get("native", ""),
                "web_crs": region.get("web_crs") or (region.get("crs") or {}).get("web", ""),
                "available_h3": region.get("available_h3_resolutions") or h3.get("available_display_resolutions", ""),
                "default_analysis_h3": region.get("default_h3_resolution") or h3.get("default_analysis_resolution", ""),
                "default_display_h3": region.get("default_display_h3_resolution") or h3.get("default_display_resolution", ""),
                "display_geometry_counts": region.get("h3_display_geometry_counts") or h3.get("display_geometry_counts", ""),
                "landing_enabled": landing.get("enabled", ""),
                "landing_badge": landing.get("badge", ""),
                "runtime_note": region.get("runtime_note", region.get("package_note", "")),
            }
        )
    add_table_sheet(wb, "Active Contract", contract_rows)

    manifest_rows = []
    for region_id, ctx in contexts.items():
        region = ctx["region"]
        manifests = region.get("manifests") or {}
        for kind in ("scenarios", "landscape", "potential", "social_acceptance", "parameter_buffers", "acceptance_registry"):
            manifest_rows.append(manifest_summary(region, kind, manifests.get(kind)))
        score = region.get("establishment_placement_score") or {}
        if score:
            manifest_rows.append(manifest_summary(region, "establishment_placement_score", score.get("manifest")))
    add_table_sheet(wb, "Manifests", manifest_rows)

    readiness_rows = []
    for region_id, ctx in contexts.items():
        region = ctx["region"]
        potential_manifest = manifest_summary(region, "potential", (region.get("manifests") or {}).get("potential"))
        readiness_rows.extend(
            [
                {
                    "region_id": region_id,
                    "area": "region_package",
                    "status": region.get("data_status", ""),
                    "review_note": "Active landing card and region package are present.",
                },
                {
                    "region_id": region_id,
                    "area": "h3_display",
                    "status": "ready_for_configured_resolutions",
                    "review_note": f"Configured H3: {join_values(region.get('available_h3_resolutions'))}; default display {region.get('default_display_h3_resolution')}.",
                },
                {
                    "region_id": region_id,
                    "area": "potential_manifest",
                    "status": potential_manifest.get("status", ""),
                    "review_note": "Bornholm still uses the dynamic R10 scaffold; Trondelag uses a placeholder manifest.",
                },
                {
                    "region_id": region_id,
                    "area": "parameter_buffer_catalog",
                    "status": "review_catalog_created",
                    "review_note": rel(catalog_paths[region_id].relative_to(ROOT)),
                },
            ]
        )
        if region_id == "bornholm":
            readiness_rows.append(
                {
                    "region_id": region_id,
                    "area": "establishment_placement_score",
                    "status": (region.get("establishment_placement_score") or {}).get("status", ""),
                    "review_note": "R9 placement score exists and is meant as soft ranking inside the user-selected potential surface.",
                }
            )
        if region_id == "trondelag":
            readiness_rows.append(
                {
                    "region_id": region_id,
                    "area": "population_buffers",
                    "status": "proxy_policy_required",
                    "review_note": "Use dissolved polygon buffers from the 250 m grid proxy, not H3 user-facing overlays.",
                }
            )
    add_table_sheet(wb, "Data Readiness", readiness_rows)

    next_rows = next_action_rows()
    add_table_sheet(wb, "Next Actions", next_rows)

    REGION_CATALOG_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(REGION_CATALOG_XLSX)


def next_action_rows() -> list[dict[str, Any]]:
    return [
        {
            "priority": "P1",
            "region_id": "bornholm",
            "action": "Audit every active registry layer against the imported R9 package and v1 source inventory.",
            "reason": "The R9 region package is in place, but the clean R9 PEY parameter runtime is not rebuilt.",
        },
        {
            "priority": "P1",
            "region_id": "bornholm",
            "action": "Decide whether Bornholm PEY should be true R9-native or labelled as compatibility mode.",
            "reason": "Potential manifest is still dynamic_res10_scaffold while display/landscape are R9/R8.",
        },
        {
            "priority": "P1",
            "region_id": "trondelag",
            "action": "Keep settlement/population buffers as dissolved 250 m grid-proxy polygons.",
            "reason": "Root AGENTS instructions forbid showing Trondelag population as H3 buffer overlays.",
        },
        {
            "priority": "P2",
            "region_id": "both",
            "action": "Add validation that each region parameter catalog has source layer IDs, CRS and operation semantics.",
            "reason": "The catalog should become a testable contract before the app consumes it.",
        },
        {
            "priority": "P2",
            "region_id": "both",
            "action": "Wire region packages to parameter_buffers.json after review.",
            "reason": "The current catalogs are review contracts, not runtime contracts.",
        },
    ]


def export_parameter_workbook(contexts: dict[str, dict[str, Any]], catalog_paths: dict[str, Path]) -> None:
    wb = make_workbook()
    add_readme_sheet(
        wb,
        "README",
        [
            ("Purpose", "Human review workbook for per-region parameter and buffer catalogs."),
            ("Generated", REPORT_DATE),
            ("Catalog JSON files", [rel(path.relative_to(ROOT)) for path in catalog_paths.values()]),
            ("Status", "Review catalog. Not yet a runtime app contract."),
            ("Important Trondelag rule", "Population buffers must be dissolved polygons from the 250 m proxy in EPSG:25832."),
            ("Important Bornholm rule", "Bornholm needs R9-native PEY/runtime review before the parameter buffers can be trusted as final."),
        ],
    )

    group_headers = [
        "region_id",
        "group_id",
        "group_label",
        "analysis_kind",
        "analysis_label",
        "default_m",
        "min_m",
        "max_m",
        "step_m",
        "blend_default",
        "source_layer_count",
        "source_layer_ids",
        "readiness",
        "source_registry",
        "interpretation",
    ]
    add_table_sheet(wb, "Group Defaults", build_group_records(contexts), headers=group_headers)

    layer_headers = [
        "region_id",
        "group_id",
        "layer_id",
        "label",
        "layer_key",
        "filter_field",
        "filter_value",
        "filter_mode",
        "readiness",
        "note",
    ]
    add_table_sheet(wb, "Source Layers", build_layer_records(contexts), headers=layer_headers)

    app_rows = solar_app_defaults() + wind_app_defaults()
    app_headers = [
        "technology",
        "default_profile",
        "control_id",
        "group_id",
        "label",
        "default_active",
        "default_layer_ids",
        "default_value",
        "spec_default_value",
        "min_value",
        "max_value",
        "step",
        "operation",
        "source",
        "notes",
    ]
    add_table_sheet(wb, "App Defaults", app_rows, headers=app_headers)

    readiness_rows = []
    for region_id, path in catalog_paths.items():
        region = contexts[region_id]["region"]
        readiness_rows.extend(
            [
                {
                    "region_id": region_id,
                    "area": "catalog_file",
                    "status": "created",
                    "path": rel(path.relative_to(ROOT)),
                    "review_note": "Source for this Excel sheet.",
                },
                {
                    "region_id": region_id,
                    "area": "native_crs",
                    "status": region.get("native_crs", ""),
                    "path": rel(contexts[region_id]["region_path"].relative_to(ROOT)),
                    "review_note": "Use this CRS for buffer and area geometry.",
                },
            ]
        )
        readiness_rows.append(
            {
                "region_id": region_id,
                "area": "runtime_status",
                "status": "review_not_runtime_contract",
                "path": rel(path.relative_to(ROOT)),
                "review_note": "Wire into app only after source and semantics review.",
            }
        )
    add_table_sheet(wb, "Readiness Matrix", readiness_rows)
    add_table_sheet(wb, "Next Actions", next_action_rows())

    PARAMETER_CATALOG_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(PARAMETER_CATALOG_XLSX)


def main() -> None:
    contexts = load_region_context()
    catalog_paths = build_region_catalogs(contexts)
    export_region_workbook(contexts, catalog_paths)
    export_parameter_workbook(contexts, catalog_paths)
    print(f"Wrote {rel(REGION_CATALOG_XLSX.relative_to(ROOT))}")
    print(f"Wrote {rel(PARAMETER_CATALOG_XLSX.relative_to(ROOT))}")
    for region_id, path in catalog_paths.items():
        print(f"Wrote {region_id}: {rel(path.relative_to(ROOT))}")


if __name__ == "__main__":
    main()
